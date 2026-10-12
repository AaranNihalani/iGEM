"""
max_peaks_plots.py -- MB / Fc / ratio-vs-time plots for Max's two SWV time
courses in data/max_data/, plus a matching signal-vs-time plot for the BSA
control, and a peak-data export (xlsx + csv), run via export_peak_data.py.

Data identity (cross-checked against data/max_data/.../mmp9_20nM_MB_Fc.zip's
own README and peaks workbook, and independently re-derived from the raw
.xlsx scans here -- both agree exactly):

  - "MMP-9 20 nM"  (folder name HAS a space)  = CLEAVABLE probe electrode.
  - "MMP-9 20nM"   (folder name has NO space) = NON-CLEAVABLE probe electrode
    (the stable reference surface; Fc is present but the probe is not cut).
  - Both are single-run time courses, scan index -> time = (index-1) x 5 min,
    EXCEPT cleavable scan 2, which Max confirms was taken at t=8 min, not 5
    (see TIME_OVERRIDES). Cleavable #2 was originally an empty .xlsx export;
    replaced 2026-10-12 with the real export from Max's own archive, which
    is an ordinary, well-behaved scan (MB 9.67 nA @ -244 mV, Fc 8.45 nA @
    290 mV -- fits cleanly between scans #1 and #3).

Scan-1 validity, per Max's direct read of the traces (2026-10-11), which
overrides the earlier blanket "scan 1 is always an outlier" assumption:
  - Cleavable scan 1: NOT clipped for either MB or Fc -- both included.
  - Non-cleavable scan 1: MB genuinely clipped (sweep starts at -200 mV
    instead of -250/-260 mV) -- MB excluded; Fc at scan 1 is fine -- Fc
    included. The ratio at this scan needs both channels, so it is omitted.

Signal axis: % of the first included scan for that channel (that scan's own
value = 100%), not raw nA, per Max's request. The ratio panel is left as a
plain MB/Fc ratio -- it is already a relative quantity, not a current.

BSA control: shows a real, reproducible feature, but centred near 0-100 mV
-- clear of both the MB window (-260 to -100 mV) and the Fc window
(+150 to +400 mV) used for the probe electrodes, so it is not MB or Fc. It
gets the same signal (%) vs time treatment, tracking that real local peak
in a -50 to +150 mV window, not an MB or Fc label.

Trend line: plain straight-line (linear) interpolation between the real,
included points -- no curve-fitting, nothing estimated that isn't actually
in the data.
"""
from pathlib import Path

import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

DATA = Path(__file__).parents[2] / "data" / "max_data"
OUT = Path(__file__).parent / "max_peaks_figures"
OUT.mkdir(exist_ok=True)

MB_WINDOW = (-0.260, -0.100)   # V
FC_WINDOW = (0.150, 0.400)     # V
BSA_WINDOW = (-0.050, 0.150)   # V -- the real, reproducible BSA feature

TIME_OVERRIDES = {("MMP-9 20 nM", 2): 8}  # (folder, scan index) -> minutes

GREEN, ORANGE, PURPLE, INK, MUTE, GRID, GREY = (
    "#15876A", "#eb6834", "#4a3aa7", "#173B34", "#8A9892", "#E3E8E1", "#B9C2BC",
)

plt.rcParams.update({
    "font.family": "STIXGeneral", "mathtext.fontset": "stix",
    "font.size": 10, "axes.titlesize": 11, "axes.labelsize": 10,
    "legend.fontsize": 9, "xtick.labelsize": 9, "ytick.labelsize": 9,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.linewidth": 0.8, "lines.linewidth": 2.0,
    "figure.dpi": 100, "savefig.bbox": "tight", "savefig.pad_inches": 0.12,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6,
    "axes.axisbelow": True, "text.color": INK,
    "axes.labelcolor": INK, "xtick.color": INK, "ytick.color": INK,
    "axes.edgecolor": MUTE,
})


def load_scan(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    pot = np.array([ws.cell(r, 2).value for r in range(2, ws.max_row + 1)], float)
    net = np.array([ws.cell(r, 6).value for r in range(2, ws.max_row + 1)], float)
    return pot, net


def peak_in_window(pot, net, lo, hi):
    m = (pot >= lo) & (pot <= hi)
    i = np.argmax(net[m])
    return net[m][i] * 1e9, pot[m][i] * 1000.0  # nA, mV


def scan_time(folder, i):
    return TIME_OVERRIDES.get((folder, i), (i - 1) * 5)


def gather(folder, mb_valid_scan1, fc_valid_scan1):
    """Per-scan dicts with raw MB/Fc peaks (nA) and per-channel ok flags.
    `*_valid_scan1` controls whether scan 1 counts as a trustworthy point
    for that channel (per Max's correction)."""
    rows = []
    for i in range(1, 14):
        f = DATA / folder / f"#{i}.xlsx"
        t = scan_time(folder, i)
        if not f.exists() or f.stat().st_size == 0:
            rows.append(dict(i=i, t=t, mb=None, fc=None, mb_ok=False, fc_ok=False, missing=True))
            continue
        pot, net = load_scan(f)
        mb, mb_pot = peak_in_window(pot, net, *MB_WINDOW)
        fc, fc_pot = peak_in_window(pot, net, *FC_WINDOW)
        mb_ok = (i != 1) or mb_valid_scan1
        fc_ok = (i != 1) or fc_valid_scan1
        rows.append(dict(i=i, t=t, mb=mb, mb_pot=mb_pot, fc=fc, fc_pot=fc_pot,
                          mb_ok=mb_ok, fc_ok=fc_ok, missing=False))
    return rows


def normalise(rows, key):
    """% of the first *ok* scan for that channel, computed for every row
    that has a raw value (ok or not) so excluded points can still be shown
    against the same reference. Returns the reference scan's time, or None."""
    ok_key = f"{key}_ok"
    ref_row = next((r for r in rows if r[ok_key] and r[key] is not None), None)
    ref = ref_row[key] if ref_row else None
    for r in rows:
        r[f"{key}_pct"] = (100.0 * r[key] / ref) if (ref and r[key] is not None) else None
    return ref_row["t"] if ref_row else None


def draw_trend(ax, t, y, color, label="Scan-to-scan (linear)"):
    """Plain straight-line interpolation between the real, included points.
    No curve-fitting -- nothing is estimated that isn't actually there."""
    order = np.argsort(t)
    t = np.asarray(t, float)[order]
    y = np.asarray(y, float)[order]
    if len(t) >= 2:
        ax.plot(t, y, color=color, lw=2.0, zorder=2, label=label)


def draw_channel(ax, rows, key, color, ylabel, marker):
    ok_key = f"{key}_ok"
    pct_key = f"{key}_pct"
    included = [r for r in rows if r[ok_key] and r[pct_key] is not None]
    excluded = [r for r in rows if (not r[ok_key]) and r[pct_key] is not None]

    t_ok = [r["t"] for r in included]
    y_ok = [r[pct_key] for r in included]
    draw_trend(ax, t_ok, y_ok, color)
    ax.plot(t_ok, y_ok, marker, color=color, ms=7, lw=0, zorder=3,
             markeredgecolor=INK, markeredgewidth=0.6, label="Scan (used)")
    if excluded:
        ax.plot([r["t"] for r in excluded], [r[pct_key] for r in excluded], marker, ms=8, zorder=3,
                 markerfacecolor="none", markeredgecolor=GREY, markeredgewidth=1.8,
                 label="Scan 1 (excluded: clipped sweep)")
    ax.axhline(100, color=MUTE, lw=0.8, ls=":", zorder=1)
    ax.set_xlabel("Incubation time (min)")
    ax.set_ylabel(ylabel)
    ax.set_xlim(-2, 62)


def draw_ratio(ax, rows):
    ok = [r for r in rows if r["mb_ok"] and r["fc_ok"] and r["mb"] is not None and r["fc"] is not None]
    t_ok = [r["t"] for r in ok]
    y_ok = [r["mb"] / r["fc"] for r in ok]
    draw_trend(ax, t_ok, y_ok, PURPLE)
    ax.plot(t_ok, y_ok, "D", color=PURPLE, ms=7, lw=0, zorder=3,
             markeredgecolor=INK, markeredgewidth=0.6, label="Scan (used)")
    ax.set_xlabel("Incubation time (min)")
    ax.set_ylabel("MB / Fc peak ratio")
    ax.set_xlim(-2, 62)


def plot_condition(folder, title, fname, mb_valid_scan1, fc_valid_scan1):
    rows = gather(folder, mb_valid_scan1, fc_valid_scan1)
    normalise(rows, "mb")
    normalise(rows, "fc")

    fig, axes = plt.subplots(1, 3, figsize=(13.5, 4.0))
    draw_channel(axes[0], rows, "mb", GREEN, "MB signal (% of first scan)", "o")
    draw_channel(axes[1], rows, "fc", ORANGE, "Fc signal (% of first scan)", "s")
    draw_ratio(axes[2], rows)

    axes[0].set_title("MB peak (ΔI, -260 to -100 mV)", loc="left", fontsize=10)
    axes[1].set_title("Fc peak (ΔI, +150 to +400 mV)", loc="left", fontsize=10)
    axes[2].set_title("MB / Fc ratio (both channels required)", loc="left", fontsize=10)
    axes[0].legend(frameon=False, fontsize=7.5, loc="upper right")

    missing = [r["t"] for r in rows if r["missing"]]
    note = f"  (scan at t={missing[0]} min missing: empty export)" if missing else ""
    fig.suptitle(f"{title}{note}", fontsize=13, y=1.03)
    fig.tight_layout()
    fig.savefig(OUT / fname, dpi=220)
    plt.close(fig)
    print("saved", fname)
    return rows


def gather_bsa():
    rows = []
    for i in range(1, 14):
        f = DATA / "BSA control" / f"#{i}.xlsx"
        t = scan_time("BSA control", i)
        if not f.exists() or f.stat().st_size == 0:
            rows.append(dict(i=i, t=t, sig=None, missing=True))
            continue
        pot, net = load_scan(f)
        sig, sig_pot = peak_in_window(pot, net, *BSA_WINDOW)
        rows.append(dict(i=i, t=t, sig=sig, sig_pot=sig_pot, missing=False))
    ref = rows[0]["sig"]
    for r in rows:
        r["sig_pct"] = (100.0 * r["sig"] / ref) if r["sig"] is not None else None
    return rows


def plot_bsa_signal(rows):
    """BSA control's real, reproducible feature, tracked the same way as the
    MB/Fc channels: % of the first scan, with a trend line. Not labelled MB
    or Fc -- it is neither."""
    fig, ax = plt.subplots(figsize=(5.6, 4.4))
    ok = [r for r in rows if r["sig"] is not None]
    t_ok = [r["t"] for r in ok]
    y_ok = [r["sig_pct"] for r in ok]
    draw_trend(ax, t_ok, y_ok, INK)
    ax.plot(t_ok, y_ok, "o", color=INK, ms=7, lw=0, zorder=3,
             markeredgecolor=GREY, markeredgewidth=0.6, label="Scan")
    ax.axhline(100, color=MUTE, lw=0.8, ls=":", zorder=1)
    ax.set_xlabel("Incubation time (min)")
    ax.set_ylabel("Signal (% of first scan)")
    ax.set_xlim(-2, 62)
    ax.legend(frameon=False, fontsize=8)
    ax.set_title("20 nM MMP-9 SWV Measurements, 60 Minutes -- BSA Control\n(real feature, -50 to +150 mV; not MB or Fc)", fontsize=11)
    fig.tight_layout()
    fig.savefig(OUT / "bsa_control_signal.png", dpi=220)
    plt.close(fig)
    print("saved bsa_control_signal.png")


if __name__ == "__main__":
    cleavable = plot_condition("MMP-9 20 nM", "20 nM MMP-9 SWV Measurements, 60 Minutes -- Cleavable", "cleavable_mb_fc_ratio.png",
                                mb_valid_scan1=True, fc_valid_scan1=True)
    noncleavable = plot_condition("MMP-9 20nM", "20 nM MMP-9 SWV Measurements, 60 Minutes -- Non-cleavable", "noncleavable_mb_fc_ratio.png",
                                   mb_valid_scan1=False, fc_valid_scan1=True)
    bsa = gather_bsa()
    plot_bsa_signal(bsa)
